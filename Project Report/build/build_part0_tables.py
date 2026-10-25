#!/usr/bin/env python3
"""Convert Part 0's LoF / LoT / Abbreviation tab-rows into formatted tables."""
import re
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

part0 = Document("Draft1. Part 0.docx")

# ---------- 1. harvest ----------
lof_rows, lot_rows, abbr_rows = [], [], []
mode = None
for p in part0.paragraphs:
    t = p.text.replace("\t", " | ").strip()
    if t == "LIST OF FIGURES":
        mode = "lof"; continue
    if t == "LIST OF TABLES":
        mode = "lot"; continue
    if t == "Abbreviation":
        mode = "abbr"; continue
    if mode == "lof" and t.startswith("Fig "):
        parts = [x.strip() for x in p.text.split("\t")]
        lof_rows.append([parts[0], " ".join(parts[1:-1]), parts[-1]])
    elif mode == "lot" and t.startswith("Table ") and not t.startswith("Table No"):
        parts = [x.strip() for x in p.text.split("\t")]
        lot_rows.append([parts[0], " ".join(parts[1:-1]), parts[-1]])
    elif mode == "abbr" and "\t" in p.text:
        parts = [x.strip() for x in p.text.split("\t")]
        if len(parts) == 2 and parts[0] and parts[1]:
            abbr_rows.append(parts)
print("harvested: LoF", len(lof_rows), "| LoT", len(lot_rows), "| Abbr", len(abbr_rows))
assert len(lof_rows) == 38 and len(lot_rows) == 12, "harvest mismatch"

# ---------- 2. abbreviation audit ----------
AUDIT_ADD = {
    "CSE": "Computer Science and Engineering",
    "CI": "Continuous Integration",
    "IEEE": "Institute of Electrical and Electronics Engineers",
    "LoRa": "Long Range (wireless communication)",
    "RAM": "Random Access Memory",
    "RTK": "Real-Time Kinematic (GPS correction)",
}
abbr = {a: e for a, e in abbr_rows if a != "UML"}
abbr.update(AUDIT_ADD)
abbr_final = sorted(abbr.items(), key=lambda kv: kv[0].lower())
print("audited abbreviations:", len(abbr_final), "| removed: UML | added:", len(AUDIT_ADD))

# ---------- 3. headings and regions ----------
heads = {}
for p in part0.paragraphs:
    t = p.text.replace("\t", " | ").strip()
    for key in ("LIST OF FIGURES", "LIST OF TABLES", "Abbreviation"):
        if t == key and key not in heads:
            heads[key] = p
assert len(heads) == 3
body = part0.element.body
NEXT = {"LIST OF FIGURES": ["LIST OF TABLES"],
        "LIST OF TABLES": ["Abbreviation"],
        "Abbreviation": []}

def region_children(head_el, stop_prefixes):
    out, collecting = [], False
    for child in body.iterchildren():
        if child is head_el:
            collecting = True; continue
        if collecting:
            if child.tag == qn("w:p"):
                texts = "".join(child.itertext()).strip()
                if any(texts.startswith(s) for s in stop_prefixes):
                    break
                out.append(child)
            elif child.tag == qn("w:sectPr"):
                break
    return out

to_delete = []
for key, head_p in heads.items():
    to_delete += region_children(head_p._p, NEXT[key])
print("old row paragraphs to delete:", len(to_delete))
for el in to_delete:
    el.getparent().remove(el)

# ---------- 4. tables ----------
def make_table(headers, rows, widths, font_size=10):
    tbl = part0.add_table(rows=1, cols=len(headers))
    tbl.autofit = False
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tbl._tbl.tblPr.append(borders)

    def fill_cell(cell, text, bold=False):
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(text)
        r.bold = bold
        r.font.size = Pt(font_size)
        r.font.name = "Times New Roman"

    for i, h in enumerate(headers):
        fill_cell(tbl.rows[0].cells[i], h, bold=True)
    trPr = tbl.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trPr.append(th)
    for row in rows:
        cells = tbl.add_row().cells
        for i, val in enumerate(row):
            fill_cell(cells[i], val)
    for i, w in enumerate(widths):
        for row in tbl.rows:
            row.cells[i].width = Inches(w)
    return tbl

t1 = make_table(["Figure No", "Description", "Page No"], lof_rows,
                [1.15, 3.95, 1.15])
heads["LIST OF FIGURES"]._p.addnext(t1._tbl)
t2 = make_table(["Table No", "Description", "Page No"], lot_rows,
                [1.15, 3.95, 1.15])
heads["LIST OF TABLES"]._p.addnext(t2._tbl)
t3 = make_table(["Abbreviation", "Expansion"], abbr_final, [1.5, 4.75])
heads["Abbreviation"]._p.addnext(t3._tbl)

part0.save("Draft1. Part 0.docx")
print("saved with 3 tables")
